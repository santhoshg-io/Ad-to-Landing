import io
from typing import Dict, Any, List, Optional
import pandas as pd

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    docx = None
    DOCX_AVAILABLE = False


def generate_excel_audit_report(
    res: Dict[str, Any],
    selected_row: pd.Series,
    processed_mismatches: List[Dict[str, Any]],
    strong_matches: List[Dict[str, Any]],
    recommendations: List[str],
    monthly_spend: float = 10000.0,
    compliance_data: Optional[Dict[str, Any]] = None
) -> bytes:
    output = io.BytesIO()
    match_score = res.get("match_score", 0)
    wasted_spend = monthly_spend * ((100.0 - match_score) / 100.0)
    annual_wasted = wasted_spend * 12.0
    risk_level = "Critical (<50%)" if match_score < 50 else ("Moderate (50-80%)" if match_score <= 80 else "Low (>80%)")

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_rows = [
            {"Parameter": "Audit Pair ID", "Value": str(selected_row.get("pair_id", ""))},
            {"Parameter": "Brand / Advertiser", "Value": str(selected_row.get("advertiser", ""))},
            {"Parameter": "Vertical", "Value": str(selected_row.get("vertical", ""))},
            {"Parameter": "Final Landing URL", "Value": str(selected_row.get("final_url", ""))},
            {"Parameter": "Overall Match Score", "Value": f"{match_score}%"},
            {"Parameter": "Above-The-Fold Match Score", "Value": f"{res.get('above_the_fold_score', 0)}%"},
            {"Parameter": "System Verdict", "Value": str(res.get("alignment_status", "Audited"))},
            {"Parameter": "Monthly PPC Budget ($)", "Value": f"${monthly_spend:,.2f}"},
            {"Parameter": "Est. Monthly Wasted Spend ($)", "Value": f"${wasted_spend:,.2f}"},
            {"Parameter": "Annualized Budget Exposure ($)", "Value": f"${annual_wasted:,.2f}"},
            {"Parameter": "CRO Risk Level", "Value": risk_level},
            {"Parameter": "Mismatches Detected", "Value": len(processed_mismatches)},
            {"Parameter": "Executive Audit Summary", "Value": str(res.get("score_summary", ""))}
        ]
        if compliance_data:
            summary_rows.append({"Parameter": "Regulated Category Status", "Value": "YES" if compliance_data.get("is_regulated") else "NO (Standard Commercial)"})
            summary_rows.append({"Parameter": "Compliance Risk Tier", "Value": str(compliance_data.get("risk_tier", "Standard"))})
            summary_rows.append({"Parameter": "Unsupported Claims (Compliance Flags)", "Value": compliance_data.get("unsupported_count", 0)})

        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Audit Summary", index=False)

        m_rows = []
        for idx, m in enumerate(processed_mismatches, 1):
            m_rows.append({
                "#": idx,
                "Severity": m.get("severity", "Moderate"),
                "Category": m.get("category", "General"),
                "Above The Fold": "Yes" if m.get("is_above_the_fold_issue") else "No",
                "Ad Copy Promise": m.get("ad_claim", ""),
                "Landing Page Quote": m.get("page_quote", ""),
                "Quote Verification": m.get("verification", {}).get("status", "UNCHECKED"),
                "Issue Description": m.get("issue_description", ""),
                "Concrete Suggested Fix": m.get("suggested_fix", "")
            })
        if m_rows:
            pd.DataFrame(m_rows).to_excel(writer, sheet_name="Mismatches", index=False)
        else:
            pd.DataFrame([{"Status": "Zero Mismatches Detected (100% Message Match)"}]).to_excel(writer, sheet_name="Mismatches", index=False)

        c_rows = []
        for idx, m in enumerate(strong_matches, 1):
            c_rows.append({
                "#": idx,
                "Ad Copy Promise": m.get("ad_claim", ""),
                "Landing Page Quote": m.get("page_quote", ""),
                "Notes / Context": m.get("notes", "Confirmed")
            })
        if c_rows:
            pd.DataFrame(c_rows).to_excel(writer, sheet_name="Confirmed Matches", index=False)
        else:
            pd.DataFrame([{"Status": "No specific matches recorded"}]).to_excel(writer, sheet_name="Confirmed Matches", index=False)

        r_rows = []
        for idx, rec in enumerate(recommendations, 1):
            r_rows.append({
                "Priority Step": f"Step {idx}",
                "Concrete Action Item": rec
            })
        if r_rows:
            pd.DataFrame(r_rows).to_excel(writer, sheet_name="Action Plan", index=False)
        else:
            pd.DataFrame([{"Status": "No corrective actions required"}]).to_excel(writer, sheet_name="Action Plan", index=False)

        if compliance_data and compliance_data.get("is_regulated"):
            comp_rows = [
                {"Parameter": "Regulated Industry Status", "Value": "YES (Regulated Category)"},
                {"Parameter": "Vertical Classification", "Value": str(compliance_data.get("info", {}).get("name", "Regulated Category"))},
                {"Parameter": "Governing Regulatory Bodies", "Value": str(compliance_data.get("info", {}).get("regulators", "FTC, CFPB"))},
                {"Parameter": "Statutory Authority", "Value": str(compliance_data.get("info", {}).get("statutes", "FTC Act Section 5"))},
                {"Parameter": "Google Ads Policy Exposure", "Value": str(compliance_data.get("info", {}).get("google_policy", "Misrepresentation Policy"))},
                {"Parameter": "Compliance Risk Tier", "Value": str(compliance_data.get("risk_tier", ""))},
                {"Parameter": "Account Suspension Risk", "Value": str(compliance_data.get("account_suspension_risk", ""))},
                {"Parameter": "Unsupported Claims Count", "Value": compliance_data.get("unsupported_count", 0)},
                {"Parameter": "Statutory Risk Analysis", "Value": str(compliance_data.get("info", {}).get("statutory_risk", ""))}
            ]
            pd.DataFrame(comp_rows).to_excel(writer, sheet_name="Compliance Audit", index=False)

    return output.getvalue()


def generate_docx_audit_report(
    res: Dict[str, Any],
    selected_row: pd.Series,
    processed_mismatches: List[Dict[str, Any]],
    strong_matches: List[Dict[str, Any]],
    recommendations: List[str],
    monthly_spend: float = 10000.0,
    compliance_data: Optional[Dict[str, Any]] = None
) -> bytes:
    if not DOCX_AVAILABLE or docx is None:
        return b""

    doc = docx.Document()

    pair_id = str(selected_row.get("pair_id", "Pair"))
    brand = str(selected_row.get("advertiser", "Brand"))
    vertical = str(selected_row.get("vertical", "Vertical"))
    url = str(selected_row.get("final_url", ""))
    match_score = res.get("match_score", 0)
    wasted_spend = monthly_spend * ((100.0 - match_score) / 100.0)
    annual_wasted = wasted_spend * 12.0
    risk_level = "Critical (<50%)" if match_score < 50 else ("Moderate (50-80%)" if match_score <= 80 else "Low (>80%)")

    doc.add_heading("Ad-to-Landing-Page Message Match Audit Report", 0)

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Pair ID: {pair_id}  |  Brand: {brand}  |  Vertical: {vertical}\n").bold = True
    p_meta.add_run(f"Destination URL: {url}\n")
    p_meta.add_run(f"System Verdict: {res.get('alignment_status', 'Audited')}  |  Match Score: {match_score}%  |  Above-The-Fold: {res.get('above_the_fold_score', 0)}%\n")
    p_meta.add_run(f"Monthly PPC Budget: ${monthly_spend:,.2f}  |  Est. Monthly Wasted Spend: ${wasted_spend:,.2f}  |  Risk Level: {risk_level}\n").italic = True
    if compliance_data and compliance_data.get("is_regulated"):
        p_meta.add_run(f"Regulated Industry: {compliance_data.get('info', {}).get('name', 'Regulated')}  |  Compliance Exposure: {compliance_data.get('risk_tier', '')}\n").bold = True

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(res.get("score_summary", "No executive summary provided."))

    doc.add_heading("2. Financial Risk & Wasted Spend Analysis", level=1)
    doc.add_paragraph(
        f"With a message match score of {match_score}%, an estimated ${wasted_spend:,.2f} of the ${monthly_spend:,.2f} "
        f"monthly PPC advertising budget is exposed to bounce friction and landing page disconnects. "
        f"This projects to an annualized wasted spend risk of ${annual_wasted:,.2f}."
    )

    if compliance_data and compliance_data.get("is_regulated"):
        doc.add_heading("3. Regulatory & Legal Compliance Exposure", level=1)
        doc.add_paragraph(
            f"Industry Classification: {compliance_data.get('info', {}).get('name', 'Regulated Category')}\n"
            f"Governing Authorities: {compliance_data.get('info', {}).get('regulators', 'FTC, CFPB')}\n"
            f"Statutory References: {compliance_data.get('info', {}).get('statutes', '')}\n"
            f"Google Ads Policy: {compliance_data.get('info', {}).get('google_policy', '')}\n"
            f"Exposure Tier: {compliance_data.get('risk_tier', '')} ({compliance_data.get('account_suspension_risk', '')})\n"
            f"Unsupported Ad Claims: {compliance_data.get('unsupported_count', 0)} claims detected with no landing page substantiation."
        )

    doc.add_heading(f"4. Identified Mismatches ({len(processed_mismatches)})", level=1)
    if processed_mismatches:
        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "#"
        hdr_cells[1].text = "Severity"
        hdr_cells[2].text = "Category"
        hdr_cells[3].text = "Ad Promise"
        hdr_cells[4].text = "Landing Page Quote"
        hdr_cells[5].text = "Verification"
        hdr_cells[6].text = "Concrete Fix"

        for idx, m in enumerate(processed_mismatches, 1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = str(m.get("severity", ""))
            row_cells[2].text = str(m.get("category", ""))
            row_cells[3].text = str(m.get("ad_claim", ""))
            row_cells[4].text = str(m.get("page_quote", ""))
            row_cells[5].text = str(m.get("verification", {}).get("status", ""))
            row_cells[6].text = str(m.get("suggested_fix", ""))
    else:
        doc.add_paragraph("Zero mismatches detected. Outstanding 100% message symmetry.")

    doc.add_heading(f"5. Confirmed Message Matches ({len(strong_matches)})", level=1)
    if strong_matches:
        t_matches = doc.add_table(rows=1, cols=3)
        t_matches.style = "Table Grid"
        h = t_matches.rows[0].cells
        h[0].text = "Ad Claim"
        h[1].text = "Landing Page Supporting Text"
        h[2].text = "Notes"

        for m in strong_matches:
            r = t_matches.add_row().cells
            r[0].text = str(m.get("ad_claim", ""))
            r[1].text = str(m.get("page_quote", ""))
            r[2].text = str(m.get("notes", ""))
    else:
        doc.add_paragraph("No specific confirmed matches listed.")

    doc.add_heading("6. Priority Action Recommendations", level=1)
    if recommendations:
        for idx, rec in enumerate(recommendations, 1):
            doc.add_paragraph(f"Step {idx}: {rec}")
    else:
        doc.add_paragraph("No corrective actions required.")

    doc_io = io.BytesIO()
    doc.save(doc_io)
    return doc_io.getvalue()


def generate_bulk_excel_report(bulk_df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        bulk_df.to_excel(writer, sheet_name="CRO Risk Leaderboard", index=False)
    return output.getvalue()